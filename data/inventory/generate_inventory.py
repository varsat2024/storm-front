"""Regenerate JSON inventories for raw GIS layers and labor workbook."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import shapefile

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "raw"
OUT = Path(__file__).resolve().parent


def inventory_shapefiles(data_dir: Path) -> dict:
    layers: dict = {}
    for shp in sorted(data_dir.glob("**/*.shp")):
        rel = shp.relative_to(data_dir).as_posix()
        reader = shapefile.Reader(str(shp))
        fields = [field[0] for field in reader.fields[1:]]
        layers[rel] = {
            "feature_count": len(reader),
            "geometry_type": reader.shapeTypeName,
            "fields": fields,
        }
        reader.close()
    return layers


def preview_labor_workbook(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=True)
    sheets = []
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        rows = list(worksheet.iter_rows(max_row=5, values_only=True))
        sheets.append({"name": name, "preview_rows": [list(row) for row in rows]})
    return {"sheets": sheets}


def main() -> None:
    for bundle in ("Data1", "Data2"):
        inventory = inventory_shapefiles(RAW / bundle)
        (OUT / f"{bundle}_inventory.json").write_text(
            json.dumps(inventory, indent=2), encoding="utf-8"
        )
    labor_path = ROOT / "Install_Labor_HackOHIO.xlsx"
    (OUT / "labor_workbook_preview.json").write_text(
        json.dumps(preview_labor_workbook(labor_path), indent=2, default=str),
        encoding="utf-8",
    )
    print(f"Wrote inventories under {OUT}")


if __name__ == "__main__":
    main()
